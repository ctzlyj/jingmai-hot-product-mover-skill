import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from merchant_template_import import (  # noqa: E402
    TemplateImportError,
    create_local_drafts,
    parse_template,
)


PRODUCT_HEADERS = [
    "批次编号", "来源平台", "来源商品链接", "来源商品ID", "商品标题", "来源店铺名称",
    "主图文件夹", "详情图文件夹", "视频文件", "目标类目", "目标类目ID", "主图模式", "备注",
]
SKU_HEADERS = [
    "来源商品ID", "来源SKU ID", "商家SKU（商家编码，必填）", "款式名称", "尺寸或规格",
    "其他销售属性", "完整源销售属性", "SKU图片路径", "来源当前有效价", "京东价",
    "预计到手价", "库存", "是否启用", "价格校验", "备注",
]


def write_image(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (800, 800), "white").save(path, format="JPEG")


def write_template(path: Path, product_row: list[object], sku_rows: list[list[object]]) -> None:
    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "填写说明"
    instruction.append(["测试模板"])
    products = workbook.create_sheet("商品信息")
    products.append(PRODUCT_HEADERS)
    products.append(product_row)
    skus = workbook.create_sheet("SKU明细")
    skus.append(SKU_HEADERS)
    for row in sku_rows:
        skus.append(row)
    skus.append([None] * 9 + ["=IF(I6=\"\",\"\",1)", "=IF(J6=\"\",\"\",1)", 10000, "是", "待填写", None])
    workbook.save(path)


def make_product_row(title: str = "自粘墙贴画") -> list[object]:
    return [
        "BATCH-001", "淘宝", "https://item.taobao.com/item.htm?id=123", "123", title, "测试店铺",
        "主图", "详情图", "", "现代装饰画", 35381, "多套主图", "",
    ]


def make_sku_row(
    source_sku: str,
    vendor_sku: str,
    style: str,
    size: str,
    attributes: dict[str, str],
    image_path: str,
    source_price: float,
) -> list[object]:
    return [
        "123", source_sku, vendor_sku, style, size, "", json.dumps(attributes, ensure_ascii=False),
        image_path, source_price, "=1", "=1", 10000, "是", "", "",
    ]


class MerchantTemplateImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.rules_path = ROOT / "config" / "modern-decorative-painting.rules.json"
        directory = self.enterContext(tempfile.TemporaryDirectory())
        self.reference_path = Path(directory) / "test-reference.json"
        self.reference_path.write_text(json.dumps({
            "category": {"id": 35381},
            "brands": [{"brandId": -1}],
            "componentCatalog": [],
            "media": {"defaultCategoryId": 1, "categories": [{"categoryId": 1}]},
        }), encoding="utf-8")

    def prepare_assets(self, root: Path) -> None:
        for index in range(1, 4):
            write_image(root / "主图" / f"main_{index:02d}.jpg")
        write_image(root / "详情图" / "detail_001.jpg")
        write_image(root / "SKU图" / "style-a.jpg")
        write_image(root / "SKU图" / "style-b.jpg")

    def test_import_preserves_custom_vendor_sku_and_builds_local_draft(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare_assets(root)
            workbook_path = root / "input.xlsx"
            rows = [
                make_sku_row("sku-a-20", "MERCHANT-A20", "A款", "20×30cm", {"颜色": "A款", "尺寸": "20×30cm"}, "SKU图/style-a.jpg", 10.0),
                make_sku_row("sku-a-30", "MERCHANT-A30", "A款", "30×40cm", {"颜色": "A款", "尺寸": "30×40cm"}, "SKU图/style-a.jpg", 12.0),
                make_sku_row("sku-b-20", "MERCHANT-B20", "B款", "20×30cm", {"颜色": "B款", "尺寸": "20×30cm"}, "SKU图/style-b.jpg", 11.0),
                make_sku_row("sku-b-30", "MERCHANT-B30", "B款", "30×40cm", {"颜色": "B款", "尺寸": "30×40cm"}, "SKU图/style-b.jpg", 13.0),
            ]
            write_template(workbook_path, make_product_row(), rows)

            parsed = parse_template(workbook_path)
            self.assertEqual(1, len(parsed))
            self.assertEqual(2, parsed[0]["inference"]["salesDimensionCount"])
            self.assertEqual("self_adhesive", parsed[0]["inference"]["productForm"])
            self.assertEqual("MERCHANT-A20", parsed[0]["source"]["variants"][0]["vendorSku"])

            results = create_local_drafts(
                workbook_path, root / "drafts", self.rules_path, self.reference_path,
            )
            listing = json.loads(Path(results[0]["listing"]).read_text(encoding="utf-8"))
            self.assertEqual(
                ["MERCHANT-A20", "MERCHANT-A30", "MERCHANT-B20", "MERCHANT-B30"],
                [sku["vendorSku"] for sku in listing["skus"]],
            )
            self.assertTrue(listing["product"]["includeSizeSaleAttribute"])
            self.assertTrue(results[0]["validation"]["ok"])
            audit = load_workbook(results[0]["auditWorkbook"], data_only=False, read_only=True)
            try:
                self.assertEqual("MERCHANT-A20", audit["SKU明细"]["E3"].value)
                self.assertEqual("@", audit["SKU明细"]["E3"].number_format)
                self.assertEqual("@", audit["SKU明细"]["Q3"].number_format)
                self.assertEqual("通过", audit["SKU明细"]["R3"].value)
            finally:
                audit.close()

    def test_single_source_dimension_is_not_split_into_size_dimension(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare_assets(root)
            workbook_path = root / "input.xlsx"
            rows = [
                make_sku_row("sku-1", "CUSTOM-001", "A款", "20×30cm", {"颜色分类": "A款 20×30cm"}, "SKU图/style-a.jpg", 10),
                make_sku_row("sku-2", "CUSTOM-002", "A款", "30×40cm", {"颜色分类": "A款 30×40cm"}, "SKU图/style-b.jpg", 12),
            ]
            write_template(workbook_path, make_product_row("卷轴挂画"), rows)

            parsed = parse_template(workbook_path)[0]
            self.assertEqual(1, parsed["inference"]["salesDimensionCount"])
            results = create_local_drafts(
                workbook_path, root / "drafts", self.rules_path, self.reference_path,
            )
            listing = json.loads(Path(results[0]["listing"]).read_text(encoding="utf-8"))
            self.assertFalse(listing["product"]["includeSizeSaleAttribute"])
            self.assertEqual(["A款 20×30cm", "A款 30×40cm"], [sku["styleAlias"] for sku in listing["skus"]])

    def test_duplicate_vendor_sku_stops_before_draft_creation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare_assets(root)
            workbook_path = root / "input.xlsx"
            rows = [
                make_sku_row("sku-1", "DUPLICATE", "A款", "20×30cm", {"颜色": "A款", "尺寸": "20×30cm"}, "SKU图/style-a.jpg", 10),
                make_sku_row("sku-2", "DUPLICATE", "B款", "20×30cm", {"颜色": "B款", "尺寸": "20×30cm"}, "SKU图/style-b.jpg", 11),
            ]
            write_template(workbook_path, make_product_row(), rows)

            with self.assertRaisesRegex(TemplateImportError, "商家SKU重复"):
                create_local_drafts(workbook_path, root / "drafts", self.rules_path, self.reference_path)
            self.assertFalse((root / "drafts").exists())

    def test_missing_source_price_and_material_path_are_blocking(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "主图").mkdir()
            (root / "详情图").mkdir()
            workbook_path = root / "input.xlsx"
            rows = [
                make_sku_row("sku-1", "CUSTOM-001", "A款", "20×30cm", {"颜色分类": "A款 20×30cm"}, "SKU图/missing.jpg", 0),
            ]
            write_template(workbook_path, make_product_row(), rows)

            with self.assertRaisesRegex(TemplateImportError, "来源当前有效价"):
                parse_template(workbook_path)


if __name__ == "__main__":
    unittest.main()
