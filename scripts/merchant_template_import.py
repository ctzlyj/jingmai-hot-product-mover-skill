#!/usr/bin/env python3
"""Parse the merchant-facing workbook and build local-only Jingmai drafts."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import OrderedDict
from copy import deepcopy
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import jd_product_mover as mover
import jdo_listing_mvp as jingmai


PRODUCT_SHEET = "商品信息"
SKU_SHEET = "SKU明细"
PRODUCT_HEADERS = {
    "来源平台", "来源商品链接", "来源商品ID", "商品标题", "来源店铺名称",
    "主图文件夹", "详情图文件夹", "视频文件", "目标类目", "目标类目ID", "主图模式",
}
SKU_HEADERS = {
    "来源商品ID", "来源SKU ID", "商家SKU（商家编码，必填）", "款式名称", "尺寸或规格",
    "完整源销售属性", "SKU图片路径", "来源当前有效价", "库存", "是否启用",
}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
PLATFORM_PREFIX = {"京东": "JD", "淘宝": "TB", "天猫": "TM", "jd": "JD", "taobao": "TB", "tmall": "TM"}
STICKY_KEYWORDS = ("自粘", "粘贴", "墙贴", "贴画", "贴纸", "背胶")
HANGING_KEYWORDS = ("挂轴", "卷轴", "挂画", "轴画")
SIZE_NAMES = ("尺寸", "尺码", "规格", "大小")


class TemplateImportError(RuntimeError):
    pass


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def table_rows(workbook: Any, sheet_name: str, required_headers: set[str], key_headers: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise TemplateImportError(f"模板缺少工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(iterator)]
    except StopIteration as exc:
        raise TemplateImportError(f"工作表为空：{sheet_name}") from exc
    missing = sorted(required_headers - set(headers))
    if missing:
        raise TemplateImportError(f"{sheet_name}缺少表头：{missing}")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(iterator, start=2):
        record = {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header}
        if not any(text(record.get(header)) for header in key_headers):
            continue
        record["_row"] = row_number
        rows.append(record)
    return rows


def positive_price(value: Any, row_number: int) -> str:
    try:
        price = Decimal(text(value))
    except InvalidOperation as exc:
        raise TemplateImportError(f"SKU明细第{row_number}行来源当前有效价不是有效数字") from exc
    if price <= 0:
        raise TemplateImportError(f"SKU明细第{row_number}行来源当前有效价必须大于0")
    return str(price.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def non_negative_stock(value: Any, row_number: int) -> int:
    rendered = text(value) or "10000"
    try:
        number = Decimal(rendered)
    except InvalidOperation as exc:
        raise TemplateImportError(f"SKU明细第{row_number}行库存不是有效整数") from exc
    if number < 0 or number != number.to_integral_value():
        raise TemplateImportError(f"SKU明细第{row_number}行库存必须为非负整数")
    return int(number)


def enabled_value(value: Any, row_number: int) -> bool:
    rendered = text(value).lower()
    if rendered in {"是", "启用", "true", "1", "yes"}:
        return True
    if rendered in {"否", "禁用", "false", "0", "no"}:
        return False
    raise TemplateImportError(f"SKU明细第{row_number}行是否启用只能填写是或否")


def source_attributes(value: Any, row_number: int) -> OrderedDict[str, str]:
    rendered = text(value)
    if not rendered:
        raise TemplateImportError(f"SKU明细第{row_number}行缺少完整源销售属性，无法判断真实销售维度")
    try:
        parsed = json.loads(rendered)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, dict):
        result = OrderedDict((text(key), text(item)) for key, item in parsed.items())
    else:
        result = OrderedDict()
        for part in re.split(r"[|；;\n]+", rendered):
            part = part.strip()
            if not part:
                continue
            match = re.match(r"^([^=:：]+)\s*[=:：]\s*(.+)$", part)
            if not match:
                raise TemplateImportError(
                    f"SKU明细第{row_number}行完整源销售属性格式错误；请用JSON或 属性=值|属性=值"
                )
            result[text(match.group(1))] = text(match.group(2))
    if not result or any(not key or not item for key, item in result.items()):
        raise TemplateImportError(f"SKU明细第{row_number}行完整源销售属性存在空名称或空值")
    return result


def resolve_path(workbook_path: Path, value: Any, label: str, *, directory: bool) -> Path:
    rendered = text(value)
    if not rendered:
        raise TemplateImportError(f"缺少{label}")
    candidate = Path(rendered)
    if not candidate.is_absolute():
        candidate = workbook_path.parent / candidate
    candidate = candidate.resolve()
    valid = candidate.is_dir() if directory else candidate.is_file()
    if not valid:
        kind = "文件夹" if directory else "文件"
        raise TemplateImportError(f"{label}{kind}不存在：{candidate}")
    return candidate


def natural_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", path.name)]


def image_files(folder: Path, label: str) -> list[Path]:
    files = sorted(
        (path.resolve() for path in folder.iterdir() if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES),
        key=natural_key,
    )
    if not files:
        raise TemplateImportError(f"{label}中没有支持的图片：{folder}")
    return files


def infer_product_form(title: str, rows: list[dict[str, Any]]) -> dict[str, str]:
    evidence = " ".join([
        title,
        *(text(row.get("款式名称")) for row in rows),
        *(text(row.get("尺寸或规格")) for row in rows),
        *(text(row.get("其他销售属性")) for row in rows),
        *(text(row.get("完整源销售属性")) for row in rows),
    ])
    if any(keyword in evidence for keyword in STICKY_KEYWORDS):
        return {"productForm": "self_adhesive", "productFormLabel": "粘贴/自粘", "vendorSkuPrefix": "ZTH"}
    if any(keyword in evidence for keyword in HANGING_KEYWORDS):
        return {"productForm": "hanging_scroll", "productFormLabel": "挂轴/卷轴", "vendorSkuPrefix": "GZH"}
    return {"productForm": "decorative_painting", "productFormLabel": "其他装饰画", "vendorSkuPrefix": "GZH"}


def variable_dimensions(all_attributes: list[OrderedDict[str, str]]) -> tuple[list[str], dict[str, list[str]]]:
    names = list(all_attributes[0])
    for attributes in all_attributes[1:]:
        if set(attributes) != set(names):
            raise TemplateImportError("完整源销售属性的属性名称不一致，无法可靠判断销售维度")
    values = {name: list(dict.fromkeys(attributes[name] for attributes in all_attributes)) for name in names}
    variable = [name for name in names if len(values[name]) > 1]
    if not variable:
        variable = names[:1]
    if len(variable) > 2:
        raise TemplateImportError(f"当前只支持最多2个真实销售属性维度，检测到：{variable}")
    if len(variable) == 2:
        size_dimensions = [name for name in variable if any(keyword in name for keyword in SIZE_NAMES)]
        if len(size_dimensions) == 1:
            size_name = size_dimensions[0]
            variable = [name for name in variable if name != size_name] + [size_name]
    return variable, values


def parse_template(workbook_path: Path) -> list[dict[str, Any]]:
    workbook_path = workbook_path.resolve()
    if not workbook_path.is_file():
        raise TemplateImportError(f"模板文件不存在：{workbook_path}")
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        products = table_rows(
            workbook, PRODUCT_SHEET, PRODUCT_HEADERS,
            ("批次编号", "来源平台", "来源商品链接", "来源商品ID", "商品标题", "来源店铺名称", "主图文件夹", "详情图文件夹", "视频文件"),
        )
        sku_rows = table_rows(
            workbook, SKU_SHEET, SKU_HEADERS,
            ("来源商品ID", "来源SKU ID", "商家SKU（商家编码，必填）", "款式名称", "尺寸或规格", "完整源销售属性", "SKU图片路径", "来源当前有效价"),
        )
    finally:
        workbook.close()
    if not products:
        raise TemplateImportError("商品信息没有已填写的商品行")
    if not sku_rows:
        raise TemplateImportError("SKU明细没有真实SKU行")

    product_ids: set[str] = set()
    source_skus: set[str] = set()
    vendor_skus: set[str] = set()
    grouped_rows: dict[str, list[dict[str, Any]]] = {}
    for row in sku_rows:
        row_number = int(row["_row"])
        product_id = text(row.get("来源商品ID"))
        source_sku = text(row.get("来源SKU ID"))
        raw_vendor_sku = row.get("商家SKU（商家编码，必填）")
        vendor_sku = text(raw_vendor_sku)
        if not product_id or not source_sku or not vendor_sku:
            raise TemplateImportError(f"SKU明细第{row_number}行缺少来源商品ID、来源SKU ID或商家SKU")
        if not isinstance(raw_vendor_sku, str):
            raise TemplateImportError(f"SKU明细第{row_number}行商家SKU必须设置为文本，避免前导零丢失")
        if source_sku in source_skus:
            raise TemplateImportError(f"来源SKU ID重复：{source_sku}")
        if vendor_sku in vendor_skus:
            raise TemplateImportError(f"商家SKU重复：{vendor_sku}")
        source_skus.add(source_sku)
        vendor_skus.add(vendor_sku)
        product_ids.add(product_id)
        row["_sourceSku"] = source_sku
        row["_vendorSku"] = vendor_sku
        row["_sourcePrice"] = positive_price(row.get("来源当前有效价"), row_number)
        row["_stock"] = non_negative_stock(row.get("库存"), row_number)
        row["_enabled"] = enabled_value(row.get("是否启用"), row_number)
        row["_attributes"] = source_attributes(row.get("完整源销售属性"), row_number)
        row["_skuImage"] = resolve_path(workbook_path, row.get("SKU图片路径"), f"SKU明细第{row_number}行SKU图片", directory=False)
        grouped_rows.setdefault(product_id, []).append(row)

    parsed_products: list[dict[str, Any]] = []
    seen_products: set[str] = set()
    for product in products:
        row_number = int(product["_row"])
        product_id = text(product.get("来源商品ID"))
        if not product_id:
            raise TemplateImportError(f"商品信息第{row_number}行缺少来源商品ID")
        if product_id in seen_products:
            raise TemplateImportError(f"商品信息来源商品ID重复：{product_id}")
        seen_products.add(product_id)
        rows = grouped_rows.get(product_id) or []
        if not rows:
            raise TemplateImportError(f"商品{product_id}没有对应的SKU明细")
        platform = text(product.get("来源平台"))
        source_url = text(product.get("来源商品链接"))
        title = text(product.get("商品标题"))
        if platform not in PLATFORM_PREFIX or not source_url or not title:
            raise TemplateImportError(f"商品信息第{row_number}行来源平台、来源链接和标题均为必填")
        category_id = int(text(product.get("目标类目ID")) or 0)
        category_name = text(product.get("目标类目"))
        main_mode = text(product.get("主图模式"))
        if category_id != 35381 or category_name != "现代装饰画" or main_mode != "多套主图":
            raise TemplateImportError(f"商品{product_id}必须固定使用现代装饰画35381和多套主图")
        main_folder = resolve_path(workbook_path, product.get("主图文件夹"), f"商品{product_id}主图", directory=True)
        detail_folder = resolve_path(workbook_path, product.get("详情图文件夹"), f"商品{product_id}详情图", directory=True)
        main_paths = image_files(main_folder, f"商品{product_id}主图文件夹")
        detail_paths = image_files(detail_folder, f"商品{product_id}详情图文件夹")
        video_value = text(product.get("视频文件"))
        video_path = resolve_path(workbook_path, video_value, f"商品{product_id}视频", directory=False) if video_value else None

        attributes = [row["_attributes"] for row in rows]
        sale_names, values_by_name = variable_dimensions(attributes)
        fixed_properties = {
            name: values for name, values in values_by_name.items() if name not in sale_names
        }
        sale_properties = OrderedDict((f"template-sale-{index}", name) for index, name in enumerate(sale_names, start=1))
        variants = []
        for row in rows:
            variants.append({
                "sourceSku": row["_sourceSku"],
                "sourcePrice": row["_sourcePrice"],
                "sourceAttributes": OrderedDict((name, row["_attributes"][name]) for name in sale_names),
                "sourceFullAttributes": dict(row["_attributes"]),
                "sourceStyleLabel": text(row.get("款式名称")),
                "sourceSizeLabel": text(row.get("尺寸或规格")),
                "vendorSku": row["_vendorSku"],
                "stock": row["_stock"],
                "enabled": row["_enabled"],
                "skuImage": str(row["_skuImage"]),
            })
        inference = {
            **infer_product_form(title, rows),
            "salesDimensionCount": len(sale_names),
            "salesDimensionNames": sale_names,
            "basis": "完整源销售属性中跨真实SKU变化的属性",
        }
        source = {
            "schemaVersion": "merchant-template/v1",
            "sourcePlatform": platform.lower() if platform.isascii() else {"京东": "jd", "淘宝": "taobao", "天猫": "tmall"}[platform],
            "productKeyPrefix": PLATFORM_PREFIX[platform],
            "collectedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "sourceUrl": source_url,
            "sourceInputUrl": source_url,
            "sourceSku": product_id,
            "sourceSpu": product_id,
            "sourceShopName": text(product.get("来源店铺名称")),
            "sourceCategoryIds": [35381],
            "categoryIds": [35381],
            "title": title,
            "fullTitle": title,
            "sourceBrandId": "",
            "sourceBrandName": "",
            "saleProperties": sale_properties,
            "sourceFixedSkuProperties": fixed_properties,
            "mainImages": [str(path) for path in main_paths],
            "detailImages": [str(path) for path in detail_paths],
            "video": ({"localPath": str(video_path)} if video_path else None),
            "variants": variants,
            "templateImport": {
                "workbook": str(workbook_path),
                "batchId": text(product.get("批次编号")),
                "note": text(product.get("备注")),
                "inference": inference,
            },
        }
        parsed_products.append({"source": source, "inference": inference})
    extra_product_ids = product_ids - seen_products
    if extra_product_ids:
        raise TemplateImportError(f"SKU明细存在商品信息中未定义的来源商品ID：{sorted(extra_product_ids)}")
    return parsed_products


def copy_materials(source: dict[str, Any], folder: Path) -> dict[str, Any]:
    result = deepcopy(source)
    for key, target_name in (("mainImages", "主图"), ("detailImages", "详情图")):
        copied: list[str] = []
        for index, raw_path in enumerate(source[key], start=1):
            path = Path(raw_path)
            target = folder / target_name / f"{index:03d}{path.suffix.lower()}"
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
            copied.append(target.relative_to(folder).as_posix())
        result[key] = copied
    sku_copies: dict[str, str] = {}
    for variant in result["variants"]:
        original = str(Path(variant["skuImage"]).resolve())
        if original not in sku_copies:
            path = Path(original)
            target = folder / "SKU图" / f"{len(sku_copies) + 1:03d}{path.suffix.lower()}"
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
            sku_copies[original] = target.relative_to(folder).as_posix()
        variant["skuImage"] = sku_copies[original]
    video = result.get("video") or {}
    if video.get("localPath"):
        source_video = Path(video["localPath"])
        target = folder / "视频" / source_video.name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_video, target)
        video["localPath"] = target.relative_to(folder).as_posix()
    return result


def pricing_verification(listing: dict[str, Any]) -> dict[str, Any]:
    pricing = listing["source"]["pricing"]
    below = Decimal(str(pricing["belowCompetitor"]))
    national = Decimal("1") - Decimal(str(pricing["nationalSubsidyRate"]))
    activity = Decimal("1") - Decimal(str(pricing["activityDiscountRate"]))
    coupon = Decimal(str(pricing["coupon"]))
    rows = []
    for sku in listing["skus"]:
        source_price = Decimal(str(sku["sourcePrice"]))
        jd_price = Decimal(str(sku["jdPrice"]))
        arrival = (jd_price * national * activity - coupon).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        target = (source_price - below).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        rows.append({
            "sourceSku": sku["sourceSku"], "vendorSku": sku["vendorSku"],
            "sourcePrice": str(source_price), "jdPrice": str(jd_price),
            "expectedArrivalPrice": str(arrival), "targetArrivalPrice": str(target),
            "ok": arrival == target,
        })
    return {"ok": all(row["ok"] for row in rows), "rows": rows}


def safe_folder_name(value: str) -> str:
    rendered = re.sub(r'[<>:"/\\|?*]', "", value).strip()
    return rendered[:60] or "template-product"


def finalize_validated_audit(path: Path, sku_count: int) -> None:
    workbook = load_workbook(path)
    sheet = workbook["SKU明细"]
    for row_number in range(3, sku_count + 3):
        sheet.cell(row_number, 5).number_format = "@"
        sheet.cell(row_number, 17).number_format = "@"
        sheet.cell(row_number, 18, "通过")
    workbook.save(path)
    workbook.close()


def create_local_drafts(
    workbook_path: Path, output_root: Path, rules_path: Path, reference_path: Path,
) -> list[dict[str, Any]]:
    parsed = parse_template(workbook_path)
    rules = mover.read_json(rules_path)
    reference = mover.read_json(reference_path)
    for item in parsed:
        if int(rules["category"]["id"]) != 35381 or int(reference["category"]["id"]) != 35381:
            raise TemplateImportError("模板导入MVP只能使用现代装饰画35381规则和实时参考")
    output_root = output_root.resolve()
    destinations = [output_root / safe_folder_name(f"{item['source']['sourceSku']}_{item['source']['title']}") for item in parsed]
    existing = [str(path) for path in destinations if path.exists()]
    if existing:
        raise TemplateImportError(f"输出目录已存在，禁止覆盖：{existing}")

    results: list[dict[str, Any]] = []
    for item, folder in zip(parsed, destinations):
        folder.mkdir(parents=True)
        source = copy_materials(item["source"], folder)
        listing = mover.build_listing(source, folder, rules, reference, int(rules["vendorSku"]["defaultStart"]))
        variants_by_sku = {variant["sourceSku"]: variant for variant in source["variants"]}
        for sku in listing["skus"]:
            variant = variants_by_sku[sku["sourceSku"]]
            sku["vendorSku"] = variant["vendorSku"]
            sku["stock"] = variant["stock"]
            sku["enabled"] = variant["enabled"]
        listing["product"]["inferredProductForm"] = item["inference"]["productForm"]
        listing["product"]["inferredProductFormLabel"] = item["inference"]["productFormLabel"]
        listing["product"]["inferredSalesDimensionCount"] = item["inference"]["salesDimensionCount"]
        listing_path = folder / "listing.json"
        errors, warnings = jingmai.validate_manifest(listing_path, listing, reference)
        validation = {"ok": not errors, "errors": errors, "warnings": warnings}
        if errors:
            raise TemplateImportError("本地草稿预检失败：" + "；".join(errors))
        pricing = pricing_verification(listing)
        if not pricing["ok"]:
            raise TemplateImportError("逐SKU价格复算失败")
        mover.write_json(folder / "manifest.json", source)
        mover.write_json(listing_path, listing)
        mover.write_json(folder / "pricing-verification.json", pricing)
        mover.write_json(folder / "local-validation.json", validation)
        audit_path = folder / "模板导入审核.xlsx"
        mover.build_audit_workbook(source, listing, audit_path)
        finalize_validated_audit(audit_path, len(listing["skus"]))
        state = {
            "stage": "LOCAL_DRAFT",
            "sourcePlatform": source["sourcePlatform"],
            "sourceSku": source["sourceSku"],
            "sourceSpu": source["sourceSpu"],
            "productId": None,
            "publishAllowed": False,
            "requires": ["AI_DUPLICATE_REVIEW", "CONTENT_COMPLIANCE_REVIEW", "PUBLISH_TO_JD"],
            "templateImport": source["templateImport"],
        }
        mover.write_json(folder / "state.json", state)
        results.append({
            "folder": str(folder),
            "manifest": str(folder / "manifest.json"),
            "listing": str(listing_path),
            "auditWorkbook": str(audit_path),
            "state": str(folder / "state.json"),
            "skuCount": len(listing["skus"]),
            "validation": validation,
        })
    return results


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="解析装饰画搬运上品标准模板并生成本地草稿")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-root", required=True, type=Path)
    parser.add_argument("--rules", type=Path, default=Path("config/modern-decorative-painting.rules.json"))
    parser.add_argument("--reference", type=Path, default=Path("config/modern-decorative-painting.reference.json"))
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        results = create_local_drafts(args.input, args.output_root, args.rules, args.reference)
    except (OSError, ValueError, TemplateImportError, mover.MoverError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 1
    print(json.dumps({"ok": True, "published": False, "results": results}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
