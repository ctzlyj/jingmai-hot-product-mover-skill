#!/usr/bin/env python3
"""Convert a local Pinduoduo export folder into a resumable Jingmai draft."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image

import jd_product_mover as mover
import jd_product_pipeline as pipeline
import jdo_listing_mvp as jingmai
from taobao_product_mvp import pricing_audit


ROOT = Path(__file__).resolve().parents[1]
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}
EXPECTED_HEADERS = (
    "上架ID", "标题", "款式", "套餐", "价格", "库存", "商品编码",
    "是否上架", "标价", "优惠券", "实际售价",
)
PRICE_HEADER_ALIASES = {"价格", "同行售价"}
DIMENSION_PATTERN = re.compile(
    r"(?:宽(?:度)?\s*\d|高(?:度)?\s*\d|\d+(?:\.\d+)?\s*(?:cm|厘米|公分|寸)|\d+\s*[*xX×]\s*\d+)",
    re.I,
)


class PddExportError(RuntimeError):
    pass


def natural_key(path: Path) -> tuple[Any, ...]:
    return tuple(
        int(part) if part.isdigit() else part.casefold()
        for part in re.split(r"(\d+)", path.name)
    )


def image_files(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return sorted(
        (path for path in folder.iterdir() if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES),
        key=natural_key,
    )


def validate_images(paths: list[Path], label: str) -> None:
    if not paths:
        raise PddExportError(f"缺少{label}")
    for path in paths:
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception as exc:
            raise PddExportError(f"{label}不可读取：{path.name}") from exc


def workbook_rows(workbook_path: Path, product_id: str) -> tuple[str, list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        if not workbook.sheetnames:
            raise PddExportError("价格工作簿没有工作表")
        sheet = workbook[workbook.sheetnames[0]]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        normalized_headers = list(headers[: len(EXPECTED_HEADERS)])
        if len(normalized_headers) >= 5 and normalized_headers[4] in PRICE_HEADER_ALIASES:
            normalized_headers[4] = "价格"
        if tuple(normalized_headers) != EXPECTED_HEADERS:
            raise PddExportError(f"价格工作簿表头不匹配：{headers[:len(EXPECTED_HEADERS)]}")

        title = ""
        rows: list[dict[str, Any]] = []
        seen_titles: set[str] = set()
        seen_ids: set[str] = set()
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = tuple(values) + (None,) * max(0, len(EXPECTED_HEADERS) - len(values))
            row_product_id, row_title, style, package, price, stock, vendor_sku = values[:7]
            has_style = style not in (None, "")
            has_price = isinstance(price, (int, float, Decimal)) and not isinstance(price, bool)
            if not has_style and not has_price:
                continue
            if not has_style or not has_price:
                raise PddExportError(f"工作簿第{row_number}行款式或价格不完整")
            if Decimal(str(price)) <= 0:
                raise PddExportError(f"工作簿第{row_number}行价格必须大于0")
            if row_product_id not in (None, ""):
                seen_ids.add(str(row_product_id).strip())
            if row_title not in (None, ""):
                seen_titles.add(str(row_title).strip())
                title = title or str(row_title).strip()
            rows.append({
                "rowNumber": row_number,
                "style": str(style).strip(),
                "package": str(package).strip() if package not in (None, "") else "",
                "sourcePrice": str(Decimal(str(price)).quantize(Decimal("0.01"))).rstrip("0").rstrip("."),
                "sourceStock": int(stock) if isinstance(stock, (int, float)) else None,
                "vendorSku": str(vendor_sku).strip() if vendor_sku not in (None, "") else "",
            })
    finally:
        workbook.close()
    if not rows:
        raise PddExportError("价格工作簿没有真实SKU行")
    if seen_ids and seen_ids != {product_id}:
        raise PddExportError(f"工作簿上架ID与目录商品ID不一致：{sorted(seen_ids)} != {product_id}")
    if len(seen_titles) > 1:
        raise PddExportError(f"工作簿存在多个商品标题：{sorted(seen_titles)}")
    if not title:
        raise PddExportError("价格工作簿缺少商品标题")
    has_package = [bool(row["package"]) for row in rows]
    if any(has_package) and not all(has_package):
        raise PddExportError("套餐维度只在部分SKU存在，无法保持真实销售属性")
    vendor_skus = [row["vendorSku"] for row in rows]
    if any(vendor_skus) and not all(vendor_skus):
        raise PddExportError("商家SKU只在部分真实SKU中填写")
    if len(set(vendor_skus)) != len(vendor_skus) and any(vendor_skus):
        raise PddExportError("商家SKU重复")
    return title, rows


def first_dimension_is_size(rows: list[dict[str, Any]]) -> bool:
    if not rows or not all(row["package"] for row in rows):
        return False
    style_rate = sum(bool(DIMENSION_PATTERN.search(row["style"])) for row in rows) / len(rows)
    package_rate = sum(bool(DIMENSION_PATTERN.search(row["package"])) for row in rows) / len(rows)
    return style_rate == 1 and package_rate == 0


def parse_export_folder(folder: Path) -> dict[str, Any]:
    source_folder = folder.resolve()
    match = re.fullmatch(r"拼多多_商品ID_(\d+)", source_folder.name)
    if not match:
        raise PddExportError(f"拼多多目录名不符合约定：{source_folder.name}")
    product_id = match.group(1)
    workbooks = sorted(source_folder.glob("*.xlsx"), key=natural_key)
    if len(workbooks) != 1:
        raise PddExportError(f"价格工作簿数量必须为1，实际为{len(workbooks)}")
    title, rows = workbook_rows(workbooks[0], product_id)
    sku_images = image_files(source_folder / "SKU")
    main_images = image_files(source_folder / "主图")
    detail_images = image_files(source_folder / "详情图")
    if len(rows) != len(sku_images):
        raise PddExportError(f"SKU行数{len(rows)}与SKU图片数{len(sku_images)}不一致")
    validate_images(sku_images, "SKU图片")
    validate_images(main_images, "主图")
    validate_images(detail_images, "详情图")

    has_package = bool(rows[0]["package"])
    swap_dimensions = has_package and first_dimension_is_size(rows)
    sale_properties = {"pdd-style": "款式"}
    if has_package:
        sale_properties["pdd-package"] = "套餐"
    variants = []
    for index, (row, image_path) in enumerate(zip(rows, sku_images), start=1):
        style_value = row["package"] if swap_dimensions else row["style"]
        package_value = row["style"] if swap_dimensions else row["package"]
        attributes = {"款式": style_value}
        if has_package:
            attributes["套餐"] = package_value
        variants.append({
            "sourceSku": f"{product_id}-SKU{index:03d}",
            "sourcePrice": row["sourcePrice"],
            "sourceListPrice": row["sourcePrice"],
            "sourceStock": row["sourceStock"],
            "vendorSku": row["vendorSku"],
            "sourceAttributes": attributes,
            "sourceRowNumber": row["rowNumber"],
            "sourceImageFile": image_path.name,
        })
    return {
        "schemaVersion": "pdd-local-export/v1",
        "sourcePlatform": "pdd",
        "sourceUrl": f"https://mobile.yangkeduo.com/goods.html?goods_id={product_id}",
        "sourceSku": product_id,
        "sourceSpu": product_id,
        "productKeyPrefix": "PDD",
        "sourceBrandName": "",
        "title": title,
        "categoryIds": [35381],
        "categoryNames": ["现代装饰画"],
        "saleProperties": sale_properties,
        "sourceFixedSkuProperties": {},
        "dimensionOrderAdjusted": swap_dimensions,
        "variants": variants,
        "sourceWorkbook": workbooks[0].name,
        "sourceMainImageFiles": [path.name for path in main_images],
        "sourceDetailImageFiles": [path.name for path in detail_images],
        "video": None,
        "collectedAt": mover.now_iso(),
    }


def copy_if_changed(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.is_file() and source.stat().st_size == destination.stat().st_size:
        if source.read_bytes() == destination.read_bytes():
            return
    shutil.copy2(source, destination)


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def image_size(path: Path) -> tuple[int, int]:
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            return image.size
    except Exception as exc:
        raise PddExportError(f"合规替换图片不可读取：{path}") from exc


def source_image_occurrences(source: dict[str, Any], source_folder: Path) -> list[dict[str, Any]]:
    occurrences: list[dict[str, Any]] = []
    for order, variant in enumerate(source["variants"], start=1):
        path = source_folder / "SKU" / variant["sourceImageFile"]
        occurrences.append({"role": "SKU", "order": order, "path": path})
    for role, key in (("主图", "sourceMainImageFiles"), ("详情图", "sourceDetailImageFiles")):
        for order, filename in enumerate(source[key], start=1):
            occurrences.append({"role": role, "order": order, "path": source_folder / role / filename})
    return occurrences


def compliance_replacements(
    source: dict[str, Any], source_folder: Path, review: dict[str, Any], review_base: Path,
) -> dict[tuple[str, int], dict[str, Any]]:
    occurrences = source_image_occurrences(source, source_folder)
    occurrence_map = {(item["role"], item["order"]): item for item in occurrences}
    replacements: dict[tuple[str, int], dict[str, Any]] = {}
    replaced_hashes: set[str] = set()
    for item in review.get("replacements") or []:
        role = str(item.get("role") or "").strip()
        order = int(item.get("order") or 0)
        key = (role, order)
        if key in replacements:
            raise PddExportError(f"合规替换位置重复声明：{role}#{order}")
        occurrence = occurrence_map.get(key)
        if occurrence is None:
            raise PddExportError(f"合规替换位置不存在：{role}#{order}")
        source_path = occurrence["path"]
        expected_filename = str(item.get("sourceFilename") or "").strip()
        if expected_filename and expected_filename != source_path.name:
            raise PddExportError(f"合规替换源文件名不匹配：{role}#{order}")
        expected_source_sha = str(item.get("sourceSha256") or "").lower()
        actual_source_sha = file_sha256(source_path)
        if not expected_source_sha or expected_source_sha != actual_source_sha:
            raise PddExportError(f"合规替换源哈希不匹配：{role}#{order}")
        replacement_path = Path(str(item.get("replacementPath") or ""))
        if not replacement_path.is_absolute():
            replacement_path = (review_base / replacement_path).resolve()
        if not replacement_path.is_file():
            raise PddExportError(f"合规替换副本不存在：{replacement_path}")
        expected_replacement_sha = str(item.get("replacementSha256") or "").lower()
        actual_replacement_sha = file_sha256(replacement_path)
        if expected_replacement_sha and expected_replacement_sha != actual_replacement_sha:
            raise PddExportError(f"合规替换副本哈希不匹配：{role}#{order}")
        if source_path.suffix.lower() != replacement_path.suffix.lower():
            raise PddExportError(f"合规替换格式必须保持不变：{role}#{order}")
        if image_size(source_path) != image_size(replacement_path) and not item.get("allowDimensionChange"):
            raise PddExportError(f"合规替换尺寸必须保持不变：{role}#{order}")
        if not str(item.get("reason") or "").strip():
            raise PddExportError(f"合规替换缺少理由：{role}#{order}")
        replacements[key] = {
            **item,
            "replacementPath": str(replacement_path),
            "sourceSha256": actual_source_sha,
            "replacementSha256": actual_replacement_sha,
        }
        replaced_hashes.add(actual_source_sha)

    undeclared = []
    for occurrence in occurrences:
        source_sha = file_sha256(occurrence["path"])
        key = (occurrence["role"], occurrence["order"])
        if source_sha in replaced_hashes and key not in replacements:
            undeclared.append(f"{occurrence['role']}#{occurrence['order']}")
    if undeclared:
        raise PddExportError("相同风险原图仍有未声明位置：" + ", ".join(undeclared))
    return replacements


def materialize_local_export(
    source: dict[str, Any], source_folder: Path, output_root: Path,
    compliance_review: dict[str, Any] | None = None, compliance_base: Path | None = None,
) -> Path:
    compliance_review = compliance_review or {}
    title_override = str(compliance_review.get("titleOverride") or "").strip()
    if title_override:
        source = deepcopy(source)
        source["sourceOriginalTitle"] = source["title"]
        source["title"] = title_override
    safe_title = re.sub(r'[<>:"/\\|?*]', "", source["title"])[:60].strip() or source["sourceSku"]
    folder = output_root / f"{source['sourceSku']}_{safe_title}"
    central_state = output_root / "state" / f"{source['sourceSku']}.json"
    state_candidates = [central_state, *(path / "state.json" for path in output_root.glob(f"{source['sourceSku']}_*") if path.is_dir())]
    for candidate in state_candidates:
        if candidate.is_file() and int((mover.read_json(candidate).get("productId") or 0)):
            raise PddExportError(f"已有productId，禁止重新生成或创建：{candidate}")
    state_path = folder / "state.json"
    if state_path.is_file() and int((mover.read_json(state_path).get("productId") or 0)):
        raise PddExportError(f"已有productId，禁止重新生成或创建：{state_path}")

    prepared = deepcopy(source)
    replacements = compliance_replacements(
        source,
        source_folder,
        compliance_review,
        (compliance_base or Path.cwd()).resolve(),
    )
    repair_audit: list[dict[str, Any]] = []
    copy_if_changed(source_folder / source["sourceWorkbook"], folder / "源文件" / source["sourceWorkbook"])
    canonical_sku_image_by_style: dict[str, Path] = {}
    for index, variant in enumerate(prepared["variants"], start=1):
        source_image = source_folder / "SKU" / variant.pop("sourceImageFile")
        relative = Path("SKU图") / f"SKU图{index:03d}{source_image.suffix.lower()}"
        replacement = replacements.get(("SKU", index))
        if replacement:
            copy_if_changed(source_image, folder / "原图归档" / relative)
            copy_if_changed(Path(replacement["replacementPath"]), folder / relative)
            repair_audit.append({
                **replacement,
                "destination": relative.as_posix(),
                "archive": (Path("原图归档") / relative).as_posix(),
            })
        else:
            copy_if_changed(source_image, folder / relative)
        style_key = str(next(iter(variant.get("sourceAttributes", {}).values()), index))
        canonical = canonical_sku_image_by_style.setdefault(style_key, relative)
        variant["skuImage"] = canonical.as_posix()

    prepared["mainImages"] = []
    for index, filename in enumerate(prepared.pop("sourceMainImageFiles"), start=1):
        source_image = source_folder / "主图" / filename
        relative = Path("主图") / f"主图{index:02d}{source_image.suffix.lower()}"
        replacement = replacements.get(("主图", index))
        if replacement:
            copy_if_changed(source_image, folder / "原图归档" / relative)
            copy_if_changed(Path(replacement["replacementPath"]), folder / relative)
            repair_audit.append({
                **replacement,
                "destination": relative.as_posix(),
                "archive": (Path("原图归档") / relative).as_posix(),
            })
        else:
            copy_if_changed(source_image, folder / relative)
        prepared["mainImages"].append(relative.as_posix())

    prepared["detailImages"] = []
    for index, filename in enumerate(prepared.pop("sourceDetailImageFiles"), start=1):
        source_image = source_folder / "详情图" / filename
        relative = Path("详情图") / f"详情图{index:03d}{source_image.suffix.lower()}"
        replacement = replacements.get(("详情图", index))
        if replacement:
            copy_if_changed(source_image, folder / "原图归档" / relative)
            copy_if_changed(Path(replacement["replacementPath"]), folder / relative)
            repair_audit.append({
                **replacement,
                "destination": relative.as_posix(),
                "archive": (Path("原图归档") / relative).as_posix(),
            })
        else:
            copy_if_changed(source_image, folder / relative)
        prepared["detailImages"].append(relative.as_posix())
    prepared["complianceRepairs"] = repair_audit
    prepared["validation"] = {"status": "passed", "errors": [], "warnings": []}
    mover.write_json(folder / "manifest.json", prepared)
    return folder


def require_review(path: Path, source_sku: str, accepted: set[str], label: str) -> dict[str, Any]:
    if not path.is_file():
        raise PddExportError(f"缺少{label}文件：{path}")
    review = mover.read_json(path)
    if str(review.get("sourceSku") or "") != source_sku:
        raise PddExportError(f"{label}的sourceSku不匹配")
    decision = str(review.get("decision") or review.get("status") or "").upper()
    if decision not in accepted:
        raise PddExportError(f"{label}未通过：{decision or '空'}")
    if not str(review.get("reason") or "").strip():
        raise PddExportError(f"{label}缺少理由")
    return review


def harmonize_sku_main_images(listing: dict[str, Any]) -> None:
    style_by_sku = {
        str(sku.get("skuKey") or ""): str(sku.get("styleAlias") or "")
        for sku in listing.get("skus") or []
    }
    canonical_by_style: dict[str, str] = {}
    for image in listing.get("images") or []:
        if image.get("role") != "sku_main":
            continue
        style = style_by_sku.get(str(image.get("skuKey") or ""), "")
        local_path = str(image.get("localPath") or "")
        if not style or not local_path:
            continue
        canonical = canonical_by_style.setdefault(style, local_path)
        image["localPath"] = canonical


def create_local_draft(args: argparse.Namespace) -> dict[str, Any]:
    source = parse_export_folder(args.source_folder)
    source_sku = str(source["sourceSku"])
    category_review = require_review(
        args.category_review, source_sku, {"PROCEED"}, "目标类目审核",
    )
    if int(category_review.get("categoryId") or 0) != 35381:
        raise PddExportError("目标类目审核不是现代装饰画35381")
    duplicate_review = require_review(
        args.duplicate_review, source_sku, {"PROCEED"}, "AI判重",
    )
    compliance_review = require_review(
        args.content_compliance, source_sku, {"PASS", "PASSED"}, "内容合规审核",
    )
    store_config = mover.read_json(args.store_config)
    store = getattr(args, "verified_store", None) or pipeline.verify_store_identity(store_config, args.profile)
    rules = pipeline.effective_rules(mover.read_json(args.rules), store_config)
    rules["pricing"] = {
        **rules["pricing"],
        "strategy": "pdd_arrival_price_multiplier",
        "targetArrivalMultiplier": "2",
        "belowCompetitor": "0",
        "targetArrivalPriceMeaning": "拼多多当前逐SKU实际售价的2倍",
    }
    reference = mover.read_json(args.reference)
    sku_limit = int(((reference.get("schema") or {}).get("extendMap") or {}).get("skuBindThresholdNum") or 299)
    if len(source["variants"]) > sku_limit:
        raise PddExportError(
            f"拼多多本地导出SKU数{len(source['variants'])}超过实时上限{sku_limit}，当前批次需先补完整款式组拆分"
        )

    title_override = str(compliance_review.get("titleOverride") or "").strip()
    if title_override:
        source["sourceOriginalTitle"] = source["title"]
        source["title"] = title_override
    prefix = mover.infer_prefix(source["title"], rules)
    style_start = int(getattr(args, "style_start", 0) or 0)
    if not style_start:
        style_start, _ = pipeline.next_style_start(prefix, 35381)
    folder = materialize_local_export(
        source,
        args.source_folder.resolve(),
        args.output_root.resolve(),
        compliance_review,
        args.content_compliance.parent.resolve(),
    )
    materialized_source = mover.read_json(folder / "manifest.json")
    listing = mover.build_listing(materialized_source, folder, rules, reference, style_start)
    harmonize_sku_main_images(listing)
    audit = pricing_audit(listing)
    if not audit["ok"]:
        raise PddExportError("拼多多逐SKU定价复算失败")
    errors, warnings = jingmai.validate_manifest(folder / "listing.json", listing, reference)
    if errors:
        raise PddExportError("本地预检失败：" + "; ".join(errors))

    mover.write_json(folder / "listing.json", listing)
    mover.write_json(folder / "pricing-verification.json", audit)
    mover.write_json(folder / "category-review.json", category_review)
    mover.write_json(folder / "duplicate-review.json", duplicate_review)
    mover.write_json(folder / "content-compliance.json", compliance_review)
    mover.build_audit_workbook(materialized_source, listing, folder / f"{folder.name}.xlsx")
    pipeline.save_state(
        folder / "state.json", "LOCAL_DRAFT", sourcePlatform="pdd",
        sourceSku=source_sku, sourceUrl=source["sourceUrl"], store=store,
        categoryReview=category_review, duplicateReview=duplicate_review,
        contentCompliance=compliance_review, skuCount=len(listing["skus"]),
        productId=None, warnings=warnings,
    )
    return {
        "ok": True,
        "stage": "LOCAL_DRAFT",
        "sourceSku": source_sku,
        "folder": str(folder),
        "skuCount": len(listing["skus"]),
        "sourceDimensions": len(source["saleProperties"]),
        "mainImages": len(materialized_source["mainImages"]),
        "detailImages": len(materialized_source["detailImages"]),
        "vendorSkuPrefix": listing["product"]["vendorSkuPrefix"],
        "vendorStyleStart": style_start,
        "pricingVerified": audit["ok"],
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-folder", type=Path, required=True)
    parser.add_argument("--category-review", type=Path, required=True)
    parser.add_argument("--duplicate-review", type=Path, required=True)
    parser.add_argument("--content-compliance", type=Path, required=True)
    parser.add_argument("--profile", required=True)
    parser.add_argument("--store-config", type=Path, required=True)
    parser.add_argument("--rules", type=Path, default=ROOT / "config" / "modern-decorative-painting.rules.json")
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--style-start", type=int)
    parser.add_argument("--output-root", type=Path, default=ROOT / "outputs" / "migrations")
    args = parser.parse_args()
    try:
        print(json.dumps(create_local_draft(args), ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, PddExportError, mover.MoverError, jingmai.JdoError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
