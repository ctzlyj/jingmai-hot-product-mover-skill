#!/usr/bin/env python3
"""Collect a JD product URL into reusable migration artifacts."""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import html as html_module
import io
import json
import re
import sys
import time
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from urllib.parse import urlencode, urlparse

import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from source_pricing import component_mode, component_source_plan


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "config" / "modern-decorative-painting.rules.json"
DEFAULT_REFERENCE = ROOT / "config" / "modern-decorative-painting.reference.json"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/152 Safari/537.36"
)


class MoverError(RuntimeError):
    pass


class SourceAccessDenied(MoverError):
    pass


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def source_sku_from_url(value: str) -> str:
    parsed = urlparse(html_module.unescape(value).replace("\\&", "&"))
    match = re.search(r"/(\d{8,})\.html", parsed.path)
    if not match:
        match = re.search(r"(?<!\d)(\d{8,})(?!\d)", value)
    if not match:
        raise MoverError(f"无法从链接识别京东SKU：{value}")
    return match.group(1)


def extract_javascript_object(page: str, variable_name: str) -> dict[str, Any]:
    marker = f"window.{variable_name} = ("
    marker_index = page.find(marker)
    if marker_index < 0:
        raise MoverError(f"页面缺少 {variable_name}")
    start = marker_index + len(marker)
    while start < len(page) and page[start].isspace():
        start += 1
    if start >= len(page) or page[start] not in "[{":
        raise MoverError(f"{variable_name} 不是对象或数组")
    depth = 0
    quote: str | None = None
    escaped = False
    for index in range(start, len(page)):
        character = page[index]
        if quote:
            if escaped:
                escaped = False
            elif character == "\\":
                escaped = True
            elif character == quote:
                quote = None
            continue
        if character in {'"', "'"}:
            quote = character
        elif character in "[{":
            depth += 1
        elif character in "]}":
            depth -= 1
            if depth == 0:
                raw = re.sub(r",\s*([}\]])", r"\1", page[start:index + 1])
                try:
                    return json.loads(raw)
                except json.JSONDecodeError as exc:
                    raise MoverError(f"{variable_name} JSON解析失败：{exc}") from exc
    raise MoverError(f"{variable_name} 对象不完整")


def parse_mobile_page(page: str) -> dict[str, Any]:
    item_only = extract_javascript_object(page, "_itemOnly")
    item_info = extract_javascript_object(page, "_itemInfo")
    item = item_only.get("item") or {}
    product = item_info.get("product") or {}
    price = ((item_info.get("priceFloor") or {}).get("ext") or {}).get("jdPrice")
    video_match = re.search(r'"mainVideoId"\s*:\s*"([^"]*)"', page)
    return {
        "item": item,
        "product": product,
        "price": price,
        "mainVideoId": video_match.group(1) if video_match else "",
    }


def parse_video_response(payload: str) -> dict[str, Any]:
    match = re.fullmatch(r"\s*videoUrlCB\((.*)\);?\s*", payload, flags=re.DOTALL)
    if not match:
        raise MoverError("视频播放地址接口返回格式异常")
    response = json.loads(match.group(1))
    if int(response.get("errcode") or 0) != 0:
        if int(response.get("errcode") or 0) in {401, 403} or re.search(r"login|unauthorized|forbidden|captcha|登录|鉴权|权限|验证码|滑块", str(response.get("errmsg") or ""), re.I):
            raise SourceAccessDenied("来源视频要求登录、权限或安全验证，停止切换工具")
        raise MoverError(f"视频播放地址查询失败：{response.get('errmsg')}")
    result = response.get("result") or {}
    variants = result.get("data") or []
    selected = next((item for item in variants if item.get("sharpness") == "high"), None)
    selected = selected or next((item for item in variants if item.get("sharpness") == "yuanhua"), None)
    selected = selected or (variants[0] if variants else None)
    if not selected or not selected.get("main_url"):
        raise MoverError("视频播放地址查询成功但未返回可下载地址")
    covers = result.get("imageUrls") or []
    return {
        "downloadUrl": str(selected["main_url"]),
        "sharpness": str(selected.get("sharpness") or ""),
        "durationSeconds": int(result.get("videoDuration") or 0),
        "sourceSizeBytes": int(result.get("videoSize") or 0),
        "coverUrl": str(covers[0]) if covers else "",
    }


def fetch_video_info(client: "HttpClient", video_id: str) -> dict[str, Any]:
    body = json.dumps(
        {"externalLoginType": 1, "callback": "videoUrlCB", "type": "1", "vid": video_id, "sceneval": "2"},
        separators=(",", ":"),
    )
    query = urlencode({
        "appid": "m_core",
        "functionId": "image_getPalyUrl",
        "jsonp": "videoUrlCB",
        "body": body,
        "appCode": "ms0ca95114",
    })
    payload = client.get_text(
        f"https://api.m.jd.com/ware/image/getPlayUrl?{query}",
        headers={"Referer": "https://item.m.jd.com/"},
    )
    return {"sourceVideoId": video_id, **parse_video_response(payload)}


def fetch_video_info_with_browser(product_url: str, video_id: str) -> dict[str, Any]:
    from browser_source_read import BrowserReadError, read_jd_video
    try:
        video = parse_video_response(read_jd_video(product_url, video_id))
        return {"sourceVideoId": video_id, "resolutionMethod": "browser-playwright", **video}
    except (BrowserReadError, MoverError, json.JSONDecodeError) as exc:
        raise MoverError(f"浏览器解析视频地址失败：{exc}") from exc


def parse_detail_urls(payload: str) -> list[str]:
    match = re.match(r"\s*showdesc\((.*)\)\s*$", payload, flags=re.DOTALL)
    if not match:
        raise MoverError("详情接口返回格式异常")
    data = json.loads(match.group(1))
    content = html_module.unescape(str(data.get("content") or ""))
    urls: list[str] = []
    # 2026-10-09: SSD detail modules embed images as CSS background-image
    # (url(...)) instead of <img src/data-lazyload>; verified on 100165442873.
    patterns = [
        r"(?:data-lazyload|src)\s*=\s*['\"]([^'\"]+)['\"]",
        r"url\(\s*['\"]?([^'\")]+?)['\"]?\s*\)",
    ]
    for pattern in patterns:
        for found in re.finditer(pattern, content, re.I):
            url = found.group(1).strip()
            if url.startswith("//"):
                url = "https:" + url
            if url.startswith(("http://", "https://")) and url not in urls:
                urls.append(url)
    return urls


def normalize_size(value: str) -> str:
    normalized = value.strip().replace("厘米", "cm")
    match = re.search(
        r"(?:宽)?\s*(\d+(?:\.\d+)?)\s*(?:cm)?\s*(?:[xX×*]|-{1,3}|[—–])\s*(?:高)?\s*(\d+(?:\.\d+)?)\s*(?:cm)?",
        normalized,
        re.I,
    )
    if not match:
        single = re.search(r"(\d+(?:\.\d+)?)\s*(寸|cm|厘米)", normalized, re.I)
        if not single:
            raise MoverError(f"无法识别尺寸：{value}")
        number = single.group(1).rstrip("0").rstrip(".") if "." in single.group(1) else single.group(1)
        unit = "cm" if single.group(2).lower() in {"cm", "厘米"} else "寸"
        return f"{number}{unit}"
    width = match.group(1).rstrip("0").rstrip(".") if "." in match.group(1) else match.group(1)
    height = match.group(2).rstrip("0").rstrip(".") if "." in match.group(2) else match.group(2)
    return f"{width}×{height}cm"


def normalize_dimension_alias(value: str) -> str:
    alias = value.strip()
    if not alias:
        raise MoverError("销售属性值为空")
    try:
        return normalize_size(alias)
    except MoverError:
        return alias


def safe_sale_alias(value: str) -> str:
    result = re.sub(r"\s+", " ", value).strip()
    while result and sum(1 if character.isascii() else 2 for character in result) > 48:
        result = result[:-1].rstrip()
    if not result:
        raise MoverError("销售属性别名为空")
    return result


def safe_filename_component(value: str, fallback: str) -> str:
    result = re.sub(r'[<>:"/\\|?*]', "", value).strip()
    return result or fallback


def disambiguate_colliding_sizes(rows: list[dict[str, Any]]) -> None:
    counts: dict[tuple[str, str], int] = {}
    for row in rows:
        key = (str(row["style"]), str(row["size"]))
        counts[key] = counts.get(key, 0) + 1
    for row in rows:
        key = (str(row["style"]), str(row["size"]))
        if counts[key] > 1:
            row["size"] = safe_sale_alias(str(row.get("sourceDimensionAlias") or row["size"]))


def size_sort_key(value: str) -> tuple[Decimal, Decimal, str]:
    if value == "单规格":
        return Decimal("0"), Decimal("0"), value
    match = re.fullmatch(r"(\d+(?:\.\d+)?)×(\d+(?:\.\d+)?)cm", value)
    if match:
        width = Decimal(match.group(1))
        height = Decimal(match.group(2))
        return max(width, height), min(width, height), value
    single = re.fullmatch(r"(\d+(?:\.\d+)?)(cm|寸)", value, re.I)
    if not single:
        raise MoverError(f"无法排序尺寸：{value}")
    length = Decimal(single.group(1))
    if single.group(2) == "寸":
        length *= Decimal("2.54")
    return length, Decimal("0"), value


def ordered_size_aliases(values: list[str]) -> list[str]:
    unique = list(dict.fromkeys(values))
    try:
        return sorted(unique, key=size_sort_key)
    except MoverError:
        return unique


def split_style_and_size(value: str, fallback_style: str) -> tuple[str, str]:
    normalized = value.strip().replace("厘米", "cm")
    dimension = re.search(
        r"(?:宽)?\s*\d+(?:\.\d+)?\s*(?:cm)?\s*(?:[xX×*]|-{1,3}|[—–])\s*(?:高)?\s*\d+(?:\.\d+)?\s*(?:cm)?",
        normalized,
        re.I,
    )
    if dimension:
        size = normalize_size(dimension.group(0))
        style = (normalized[:dimension.start()] + " " + normalized[dimension.end():]).strip()
    else:
        single = re.search(r"\d+(?:\.\d+)?\s*(?:寸|cm)", normalized, re.I)
        if not single:
            raise MoverError(f"无法从销售属性拆分款式和尺寸：{value}")
        size = normalize_size(single.group(0))
        style = (normalized[:single.start()] + " " + normalized[single.end():]).strip()
    style = re.sub(r"\s+", " ", style).strip(" -_/，,；;：:")
    return style or fallback_style, size


def variant_style_and_size(
    value: str, fallback_style: str, *, allow_single_spec: bool,
) -> tuple[str, str]:
    try:
        return split_style_and_size(value, fallback_style)
    except MoverError:
        if not allow_single_spec:
            raise
        return value.strip() or fallback_style, "单规格"


def base_title(full_title: str, sale_values: list[str]) -> str:
    result = full_title
    for value in sorted({item.strip() for item in sale_values if item.strip()}, key=len, reverse=True):
        result = result.replace(value, " ")
    result = re.sub(r"\s+", " ", result).strip(" -_/，,；;：:")
    if not result:
        raise MoverError("移除销售属性后商品标题为空")
    return result


def jd_title_units(value: str) -> int:
    return sum(1 if character.isascii() else 2 for character in value)


def listing_title(source_title: str, source_brand: str = "", max_units: int = 60) -> str:
    result = source_title
    if source_brand.strip():
        result = result.replace(source_brand.strip(), " ")
    result = re.sub(r"\s+", " ", result).strip(" -_/，,；;：:")
    if not result:
        raise MoverError("移除源品牌后商品标题为空")
    while result and jd_title_units(result) > max_units:
        result = result[:-1].rstrip()
    if not result:
        raise MoverError("商品标题无法满足京东长度限制")
    return result


def infer_style_alias(title: str) -> str:
    bracketed = re.search(r"【([^】]{2,20})】", title)
    if bracketed:
        return bracketed.group(1).strip()
    for keyword in ("喜居宝地", "家和万事兴", "一帆风顺", "背靠金山", "九鱼聚财"):
        if keyword in title:
            return keyword
    return "默认款式"


def infer_prefix(title: str, rules: dict[str, Any]) -> str:
    vendor_rules = rules["vendorSku"]
    for rule in vendor_rules.get("prefixRules") or []:
        if any(keyword in title for keyword in rule.get("keywords") or []):
            return str(rule["prefix"])
    return str(vendor_rules["defaultPrefix"])


def infer_subject(title: str, rules: dict[str, Any]) -> dict[str, str]:
    subject = rules["skuAttributes"]["subject"]
    for rule in subject.get("rules") or []:
        if any(keyword in title for keyword in rule.get("keywords") or []):
            return {"valueId": str(rule["valueId"]), "label": str(rule["label"])}
    default = subject["default"]
    return {"valueId": str(default["valueId"]), "label": str(default["label"])}


def required_sku_attribute_values(
    required_components: list[dict[str, Any]], rules: dict[str, Any],
) -> list[dict[str, str]]:
    sku_rules = rules.get("skuAttributes") or {}
    defaults = sku_rules.get("requiredDefaults") or {}
    legacy_defaults: dict[str, dict[str, str]] = {}
    for key, title in (("mounting", "装裱"), ("subject", "题材")):
        default = (sku_rules.get(key) or {}).get("default") or {}
        if default.get("valueId"):
            legacy_defaults[title] = {
                "valueId": str(default["valueId"]),
                "label": str(default.get("label") or ""),
            }
    resolved: list[dict[str, str]] = []
    missing: list[str] = []
    for component in required_components:
        code = str(component.get("componentCode") or "")
        title = str(component.get("title") or code)
        default = defaults.get(code) or defaults.get(title) or legacy_defaults.get(title)
        if not default or not default.get("valueId"):
            missing.append(title)
            continue
        resolved.append({
            "componentCode": code,
            "valueId": str(default["valueId"]),
            "label": str(default.get("label") or ""),
        })
    if missing:
        raise MoverError(f"类目存在未适配的必填SKU属性：{missing}")
    return resolved


def infer_scenes(title: str, rules: dict[str, Any]) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for rule in rules.get("sceneRules") or []:
        if rule["keyword"] in title and rule["valueId"] not in {item["valueId"] for item in result}:
            result.append({"valueId": str(rule["valueId"]), "label": str(rule["label"])})
        if len(result) == 3:
            break
    return result or [{"valueId": str(item["valueId"]), "label": str(item["label"])} for item in rules.get("defaultScenes") or []]


def target_price(source_price: str | Decimal, pricing: dict[str, Any]) -> Decimal:
    if component_mode(pricing):
        return component_source_plan(source_price, pricing).jd_price
    source = Decimal(str(source_price))
    multiplier = Decimal(str(pricing.get("targetArrivalMultiplier") or "1"))
    below = Decimal(str(pricing["belowCompetitor"]))
    coupon = Decimal(str(pricing["coupon"]))
    national = Decimal("1") - Decimal(str(pricing["nationalSubsidyRate"]))
    activity = Decimal("1") - Decimal(str(pricing["activityDiscountRate"]))
    return ((source * multiplier - below + coupon) / (national * activity)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def stable_mapping_id(product_key: str, field: str, value: str) -> str:
    return hashlib.sha256(f"{product_key}|{field}|{value}".encode()).hexdigest()[:32]


class HttpClient:
    def __init__(self, timeout: int = 45, retries: int = 4):
        self.timeout = timeout
        self.retries = retries
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "zh-CN,zh;q=0.9"})

    def get_bytes(self, url: str, headers: dict[str, str] | None = None) -> bytes:
        last_error: Exception | None = None
        for attempt in range(self.retries):
            try:
                response = self.session.get(url, timeout=self.timeout, headers=headers)
                response.raise_for_status()
                return response.content
            except requests.RequestException as exc:
                if exc.response is not None and exc.response.status_code in {401, 403}:
                    raise SourceAccessDenied("来源接口要求登录或权限，停止重试和浏览器回退") from exc
                last_error = exc
                if attempt + 1 < self.retries:
                    time.sleep(2 ** attempt)
        raise MoverError(f"请求失败：{url}：{last_error}")

    def get_text(self, url: str, headers: dict[str, str] | None = None) -> str:
        return self.get_bytes(url, headers=headers).decode("utf-8", errors="replace")

    def download(self, url: str, destination: Path) -> None:
        last_error: Exception | None = None
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(destination.suffix + ".part")
        for attempt in range(self.retries):
            try:
                head = self.session.head(url, timeout=self.timeout, allow_redirects=True)
                head.raise_for_status()
                total_size = int(head.headers.get("content-length") or 0)
                if total_size <= 0:
                    raise MoverError("下载响应缺少Content-Length")
                offset = temporary.stat().st_size if temporary.exists() else 0
                if offset > total_size:
                    temporary.unlink()
                    offset = 0
                with temporary.open("ab") as output:
                    while offset < total_size:
                        end = min(offset + 4 * 1024 * 1024 - 1, total_size - 1)
                        response = self.session.get(
                            url,
                            headers={"Range": f"bytes={offset}-{end}"},
                            timeout=self.timeout,
                        )
                        response.raise_for_status()
                        chunk = response.content
                        if not chunk:
                            raise MoverError(f"分段下载返回空内容：{offset}-{end}")
                        output.write(chunk)
                        output.flush()
                        offset += len(chunk)
                if temporary.stat().st_size != total_size:
                    raise MoverError(
                        f"下载文件大小不一致：{temporary.stat().st_size}，预期{total_size}"
                    )
                temporary.replace(destination)
                return
            except (OSError, requests.RequestException, MoverError) as exc:
                last_error = exc
                if attempt + 1 < self.retries:
                    time.sleep(2 ** attempt)
        raise MoverError(f"下载失败：{url}：{last_error}")


def jd_image_url(path: str, *, square: bool) -> str:
    if path.startswith("//"):
        return "https:" + path
    if path.startswith(("http://", "https://")):
        return path
    prefix = "https://img10.360buyimg.com/n1/s800x800_" if square else "https://img10.360buyimg.com/img/"
    return prefix + path.lstrip("/")


def save_png(content: bytes, destination: Path, *, square: bool) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        with Image.open(io.BytesIO(content)) as image:
            image.load()
            if image.mode not in {"RGB", "RGBA"}:
                image = image.convert("RGBA" if "transparency" in image.info else "RGB")
            if square:
                image.thumbnail((800, 800), Image.Resampling.LANCZOS)
                canvas = Image.new("RGB", (800, 800), "white")
                position = ((800 - image.width) // 2, (800 - image.height) // 2)
                canvas.paste(image, position, image if image.mode == "RGBA" else None)
                image = canvas
            elif image.mode == "RGBA":
                background = Image.new("RGB", image.size, "white")
                background.paste(image, mask=image)
                image = background
            image.save(destination, "PNG", optimize=True)
    except Exception as exc:
        raise MoverError(f"图片无法转换：{destination.name}：{exc}") from exc
    if destination.stat().st_size >= 3 * 1024 * 1024:
        raise MoverError(f"转换后图片仍达到3MB：{destination}")


def collect_source(url: str, client: HttpClient, max_workers: int = 8) -> dict[str, Any]:
    input_sku = source_sku_from_url(url)
    canonical_url = f"https://item.jd.com/{input_sku}.html"
    first = parse_mobile_page(client.get_text(f"https://item.m.jd.com/product/{input_sku}.html"))
    item = first["item"]
    product = first["product"]
    category_ids = [int(value) for value in str(product.get("category") or "").split(",") if value]
    variants = item.get("newColorSize") or [{"skuId": input_sku, "imagePath": product.get("imageurl")}]
    sku_ids = [str(variant.get("skuId") or "") for variant in variants]
    if not all(sku_ids) or len(set(sku_ids)) != len(sku_ids):
        raise MoverError("源商品SKU ID缺失或重复")

    def fetch_variant(sku_id: str) -> tuple[str, dict[str, Any]]:
        page = client.get_text(f"https://item.m.jd.com/product/{sku_id}.html")
        return sku_id, parse_mobile_page(page)

    pages: dict[str, dict[str, Any]] = {input_sku: first}
    with concurrent.futures.ThreadPoolExecutor(max_workers=min(max_workers, len(sku_ids))) as executor:
        futures = [executor.submit(fetch_variant, sku_id) for sku_id in sku_ids if sku_id != input_sku]
    for future in concurrent.futures.as_completed(futures):
        sku_id, parsed = future.result()
        pages[sku_id] = parsed
    prices = {sku_id: str(pages[sku_id].get("price") or "") for sku_id in sku_ids}
    image_sequences = {
        sku_id: [
            jd_image_url(str(path), square=(position == 1))
            for position, path in enumerate(
                (pages[sku_id].get("item") or {}).get("image") or [], start=1
            )
        ]
        for sku_id in sku_ids
    }
    if any(not re.fullmatch(r"\d+(?:\.\d+)?", price) for price in prices.values()):
        missing = [sku_id for sku_id, price in prices.items() if not re.fullmatch(r"\d+(?:\.\d+)?", price)]
        raise MoverError(f"逐SKU价格缺失：{missing}")
    sale_prop = {str(key): str(value) for key, value in (item.get("saleProp") or {}).items()}
    sale_values = [str(variant.get(key) or "") for variant in variants for key in sale_prop]
    title = base_title(str(product.get("skuName") or item.get("skuName") or ""), sale_values)
    # 2026-10-09: self-operated products lack popPid; realSpuId works for desc API (verified on 100165442873).
    pop_pid = str(product.get("popPid") or product.get("realSpuId") or "")
    if not pop_pid:
        raise MoverError("源商品缺少popPid，无法采集详情")
    detail_urls = parse_detail_urls(client.get_text(f"https://d.3.cn/desc/{pop_pid}"))
    video_id = str(first.get("mainVideoId") or "")
    video = None
    if video_id:
        try:
            video = {"resolutionMethod": "api", **fetch_video_info(client, video_id)}
        except SourceAccessDenied:
            raise
        except MoverError:
            video = fetch_video_info_with_browser(
                f"https://item.m.jd.com/product/{input_sku}.html",
                video_id,
            )
    else:
        # 2026-09-30: some source videos are PC-page-only (no mobile
        # mainVideoId). Probe the rendered PC page before reporting no video;
        # 10201583840642 was initially mis-collected as video-less this way.
        from browser_source_read import BrowserReadError, read_jd_pc_video
        try:
            video = {"resolutionMethod": "browser-pc-video-element",
                     **read_jd_pc_video(f"https://item.jd.com/{input_sku}.html")}
        except (BrowserReadError, MoverError, OSError):
            video = None
    return {
        "schemaVersion": "jd-mover/v1",
        "collectedAt": now_iso(),
        "sourceUrl": canonical_url,
        "sourceInputUrl": url,
        "sourceSku": input_sku,
        "sourceSpu": pop_pid,
        "sourceVenderId": str(product.get("venderId") or item.get("venderID") or ""),
        "categoryIds": category_ids,
        "title": title,
        "fullTitle": str(product.get("skuName") or ""),
        "sourceBrandId": int(product.get("brandId") or -1),
        "sourceBrandName": str(item.get("brandName") or ""),
        "saleProperties": sale_prop,
        "mainImageUrls": [jd_image_url(str(path), square=True) for path in item.get("image") or []],
        "detailImageUrls": detail_urls,
        "videoUrl": video["downloadUrl"] if video else "",
        "video": video,
        "variants": [
            {
                "sourceSku": sku_id,
                "sourcePrice": prices[sku_id],
                "sourceAttributes": {name: str(variant.get(key) or "") for key, name in sale_prop.items()},
                "sourceImageUrl": jd_image_url(str(variant.get("imagePath") or product.get("imageurl") or ""), square=True),
                "sourceImageUrls": image_sequences[sku_id],
            }
            for sku_id, variant in zip(sku_ids, variants)
        ],
    }


def materialize(source: dict[str, Any], output_root: Path, client: HttpClient) -> Path:
    safe_title = re.sub(r'[<>:"/\\|?*]', "", source["title"])[:60].strip() or source["sourceSku"]
    folder = output_root / f"{source['sourceSku']}_{safe_title}"
    from hot_replica_mode import style_image_mode

    style_layout = style_image_mode(source)
    source["styleImageMode"] = style_layout
    source["styleSceneImagesBySku"] = {}
    for name in ("主图", "SKU图", "详情图", "视频", *(["款式组原图"] if style_layout["mode"] == "series" else [])):
        (folder / name).mkdir(parents=True, exist_ok=True)

    main_paths: list[str] = []
    for index, url in enumerate(source["mainImageUrls"], start=1):
        relative = Path("主图") / f"主图{index}.png"
        destination = folder / relative
        if not destination.exists():
            save_png(client.get_bytes(url), destination, square=True)
        main_paths.append(relative.as_posix())

    style_urls = list(dict.fromkeys(variant["sourceImageUrl"] for variant in source["variants"]))
    sku_image_by_url: dict[str, str] = {}
    for index, url in enumerate(style_urls, start=1):
        relative = Path("SKU图") / f"SKU图{index}.png"
        destination = folder / relative
        if not destination.exists():
            save_png(client.get_bytes(url), destination, square=True)
        sku_image_by_url[url] = relative.as_posix()

    detail_paths: list[str] = []
    for index, url in enumerate(source["detailImageUrls"], start=1):
        relative = Path("详情图") / f"详情图{index:03d}.png"
        destination = folder / relative
        if not destination.exists():
            save_png(client.get_bytes(url), destination, square=False)
        detail_paths.append(relative.as_posix())

    video = source.get("video")
    if video:
        video_relative = Path("视频") / "主图视频.mp4"
        video_destination = folder / video_relative
        if not video_destination.exists():
            client.download(str(video["downloadUrl"]), video_destination)
        video["localPath"] = video_relative.as_posix()
        video["downloadedSizeBytes"] = video_destination.stat().st_size
        if video.get("coverUrl"):
            cover_relative = Path("视频") / "视频封面.jpg"
            cover_destination = folder / cover_relative
            if not cover_destination.exists():
                client.download(str(video["coverUrl"]), cover_destination)
            video["coverLocalPath"] = cover_relative.as_posix()

    if source.get("sourcePlatform") == "taobao" and source.get("imageQualityGate") == "strict":
        import shutil
        snapshots = {}
        for relative in dict.fromkeys([*main_paths, *detail_paths, *sku_image_by_url.values()]):
            original = (Path("source-original") / relative).as_posix()
            path = folder / original
            path.parent.mkdir(parents=True, exist_ok=True)
            if not path.exists():
                shutil.copyfile(folder / relative, path)
            snapshots[relative] = original
        source["sourceImageSnapshot"] = snapshots

    source["mainImages"] = main_paths
    source["detailImages"] = detail_paths
    for variant in source["variants"]:
        variant["skuImage"] = sku_image_by_url[variant["sourceImageUrl"]]
    if style_layout["mode"] == "series":
        for variant in source["variants"]:
            style = safe_sale_alias(str(next(iter((variant.get("sourceAttributes") or {}).values()), variant["sourceSku"])))
            style_filename = safe_filename_component(style, str(variant["sourceSku"]))
            scene_paths: dict[str, str] = {}
            for position, url in enumerate(variant.get("sourceImageUrls") or [], start=1):
                if position == 1:
                    continue
                relative = Path("款式组原图") / f"{style_filename}-pos{position}.png"
                destination = folder / relative
                if not destination.exists():
                    save_png(client.get_bytes(url), destination, square=False)
                scene_paths[str(position)] = relative.as_posix()
            existing = source["styleSceneImagesBySku"].get(variant["sourceSku"])
            if existing is not None and existing != scene_paths:
                raise MoverError(f"同一款式不同SKU的主图序列不一致：{variant['sourceSku']}")
            source["styleSceneImagesBySku"][variant["sourceSku"]] = scene_paths
    source["validation"] = {"status": "passed", "errors": [], "warnings": []}
    write_json(folder / "manifest.json", source)
    return folder


def video_handoff(folder: Path, listing: dict[str, Any]) -> dict[str, Any] | None:
    video = listing.get("video") or {}
    local_path = str(video.get("localPath") or "").strip()
    if not local_path:
        return None
    video_path = (folder / local_path).resolve()
    if not video_path.is_file():
        raise MoverError(f"源商品存在视频，但本地视频文件缺失：{video_path}")
    cover_path = None
    if str(video.get("coverLocalPath") or "").strip():
        cover_path = (folder / str(video["coverLocalPath"])).resolve()
        if not cover_path.is_file():
            raise MoverError(f"视频封面文件缺失：{cover_path}")
    return {
        "manualUploadRequired": False,
        "automaticUploadPlanned": True,
        "videoPath": str(video_path),
        "coverPath": str(cover_path) if cover_path else "",
        "durationSeconds": int(video.get("durationSeconds") or 0),
        "sizeBytes": video_path.stat().st_size,
        "sourceVideoId": str(video.get("sourceVideoId") or ""),
        "message": f"该商品含主图视频，已保存到：{video_path}。商品发布后将自动上传并绑定全部SKU。",
    }


def expected_source_category_id(rules: dict[str, Any]) -> int:
    category = rules["category"]
    return int(category.get("sourceId") or category["id"])


def build_listing(
    source: dict[str, Any], folder: Path, rules: dict[str, Any],
    reference: dict[str, Any], style_start: int,
) -> dict[str, Any]:
    expected_category = expected_source_category_id(rules)
    review = source.get("categoryReview") or {}
    reviewed_mapping = (
        source.get("sourcePlatform") == "taobao" and review.get("decision") == "PROCEED"
        and bool(str(review.get("reason") or "").strip())
        and str(review.get("sourceItemId") or "") == str(source.get("sourceParentSku") or source["sourceSku"])
        and int(review.get("targetCategoryId") or 0) == int(rules["category"]["id"])
        and int((reference.get("category") or {}).get("id") or 0) == int(rules["category"]["id"])
    )
    if not reviewed_mapping and (not source.get("categoryIds") or int(source["categoryIds"][-1]) != expected_category):
        raise MoverError(f"源商品类目与规则不一致，预期来源类目{expected_category}：{source.get('categoryIds')}")
    product_key = f"{str(source.get('productKeyPrefix') or 'JD').upper()}-{source['sourceSpu']}"
    source_title = source["title"]
    title = listing_title(source_title, str(source.get("sourceBrandName") or ""))
    prefix = infer_prefix(source_title, rules)
    scenes = infer_scenes(source_title, rules)
    sku_rules = rules.get("skuAttributes") or {}
    subject = infer_subject(source_title, rules) if sku_rules.get("subject") else {"valueId": "", "label": ""}
    mounting = (sku_rules.get("mounting") or {}).get("default") or {"valueId": "", "label": ""}
    sale_names = list(source.get("saleProperties", {}).values())
    if len(sale_names) > 2:
        raise MoverError("当前发布模型最多表达两个销售维度，不能合并或丢弃源属性")
    has_style_dimension = len(sale_names) >= 2 and sale_names[0] not in {"尺寸", "规格"}
    fallback_style = infer_style_alias(source_title)
    component_catalog = reference.get("componentCatalog") or []
    sale_components = [
        item for item in component_catalog
        if item.get("path") == "root/SKUS/VenderCategoryValue"
        and str(item.get("componentCode") or "").startswith("sale-prop-")
    ]
    style_sale_component = next(
        (item for item in sale_components if str(item.get("title") or "") in {"颜色", "款式", "规格"}),
        sale_components[0] if sale_components else None,
    )
    size_sale_component = next(
        (item for item in sale_components if str(item.get("title") or "") in {"尺寸", "尺码"}),
        sale_components[1] if len(sale_components) > 1 else None,
    )
    sale_attribute_defaults = rules.get("saleAttributes") or {}
    required_sku_attributes = [
        item for item in component_catalog
        if item.get("path") == "root/SKUS/SkuAttr" and item.get("required")
    ]
    resolved_sku_attributes = required_sku_attribute_values(required_sku_attributes, rules)

    prepared: list[dict[str, Any]] = []
    for variant in source["variants"]:
        attributes = variant.get("sourceAttributes") or {}
        values = list(attributes.values())
        if has_style_dimension:
            style = values[0]
            vendor_style = style
            size_source = " ".join(values[1:])
            size = normalize_dimension_alias(size_source)
        elif len(sale_names) == 1 and values:
            vendor_style, size = variant_style_and_size(
                values[0], fallback_style, allow_single_spec=len(sale_names) == 1,
            )
            style = values[0].strip()
        else:
            style = fallback_style
            vendor_style = style
            size_source = next(
                (value for name, value in attributes.items() if name in {"尺寸", "规格"}),
                values[-1] if values else "",
            )
            size = normalize_size(size_source)
        prepared.append({
            **variant,
            "style": safe_sale_alias(style),
            "vendorStyle": safe_sale_alias(vendor_style),
            "size": safe_sale_alias(size),
            "sourceDimensionAlias": size_source if has_style_dimension else size,
        })
    disambiguate_colliding_sizes(prepared)
    style_names = list(dict.fromkeys(item["style"] for item in prepared))
    size_names = ordered_size_aliases([item["size"] for item in prepared])
    vendor_style_names = list(dict.fromkeys(item["vendorStyle"] for item in prepared))
    style_orders = {name: index for index, name in enumerate(style_names, start=1)}
    size_orders = {name: index for index, name in enumerate(size_names, start=1)}
    vendor_style_orders = {name: index for index, name in enumerate(vendor_style_names, start=1)}
    vendor_size_orders: dict[tuple[str, str], int] = {}
    for vendor_style in vendor_style_names:
        sizes = ordered_size_aliases([
            item["size"] for item in prepared if item["vendorStyle"] == vendor_style
        ])
        for order, size in enumerate(sizes, start=1):
            vendor_size_orders[(vendor_style, size)] = order

    defaults = rules["productDefaults"]
    category_attributes = [*rules["categoryAttributes"]]
    if any(item.get("componentCode") == "prop-product-243294" for item in component_catalog):
        category_attributes.append({
            "componentCode": "prop-product-243294",
            "attrName": "适用场景",
            "inputType": "multiSelect",
            "required": True,
            "valueIds": [item["valueId"] for item in scenes],
            "valueLabels": [item["label"] for item in scenes],
        })

    listing: dict[str, Any] = {
        "schemaVersion": 1,
        "scene": "pop",
        "source": {
            "url": source["sourceUrl"],
            "sku": source["sourceSku"],
            "spu": source["sourceSpu"],
            "collectedAt": source["collectedAt"],
            "pricing": rules["pricing"],
        },
        "category": rules["category"],
        "product": {
            "productKey": product_key,
            "title": title,
            "brandId": defaults["brandId"],
            "brandName": defaults["brandName"],
            "model": str(defaults.get("model") or ""),
            "itemNumber": source["sourceSpu"],
            "packageLengthMm": defaults.get("packageLengthMm", ""),
            "packageWidthMm": defaults.get("packageWidthMm", ""),
            "packageHeightMm": defaults.get("packageHeightMm", ""),
            "weightKg": defaults.get("weightKg", ""),
            "sevenDayReturnCode": defaults["sevenDayReturnCode"],
            "sevenDayReturnLabel": defaults["sevenDayReturnLabel"],
            "saleUnit": "",
            "freightTemplateId": defaults["freightTemplateId"],
            "freightTemplateLabel": defaults["freightTemplateLabel"],
            "mainImageMode": defaults["mainImageMode"],
            "styleImageMode": (source.get("styleImageMode") or {}).get("mode", "shared"),
            "includeSizeSaleAttribute": len(sale_names) >= 2,
            "detailHtml": "",
            "defaultMediaCategoryId": reference["media"]["defaultCategoryId"],
            "vendorSkuPrefix": prefix,
            "vendorStyleStart": style_start,
            "saleAttributeComponents": {
                "style": ({
                    "componentCode": style_sale_component["componentCode"],
                    "attrName": style_sale_component.get("title") or "颜色",
                    **(sale_attribute_defaults.get("style") or {}),
                } if style_sale_component else None),
                "size": ({
                    "componentCode": size_sale_component["componentCode"],
                    "attrName": size_sale_component.get("title") or "尺寸",
                    **(sale_attribute_defaults.get("size") or {}),
                } if size_sale_component else None),
            },
        },
        "skus": [],
        "categoryAttributes": category_attributes,
        "video": source.get("video"),
        "images": [],
    }
    for index, item in enumerate(prepared, start=1):
        style_order = style_orders[item["style"]]
        size_order = size_orders[item["size"]]
        sku_key = f"{product_key}-SKU-{index:03d}"
        listing["skus"].append({
            "skuKey": sku_key,
            "productKey": product_key,
            "vendorSku": str(item.get("vendorSku") or "").strip(),
            "jdPrice": float(target_price(item["sourcePrice"], rules["pricing"])),
            "stock": int(item.get("stock", defaults["stock"])),
            "shortTitle": "",
            "enabled": bool(item.get("enabled", True)),
            "mountingValueId": str(mounting["valueId"]),
            "mountingLabel": mounting["label"],
            "subjectValueId": str(subject["valueId"]),
            "subjectLabel": subject["label"],
            "styleAlias": item["style"],
            "styleMappingId": stable_mapping_id(product_key, "style", item["style"]),
            "styleOrder": style_order,
            "vendorStyleAlias": item["vendorStyle"],
            "vendorStyleOrder": vendor_style_orders[item["vendorStyle"]],
            "sizeAlias": item["size"],
            "sizeMappingId": stable_mapping_id(product_key, "size", item["size"]),
            "sizeOrder": size_order,
            "vendorSizeOrder": vendor_size_orders[(item["vendorStyle"], item["size"])],
            "sourceSku": item["sourceSku"],
            "sourcePrice": float(Decimal(item["sourcePrice"])),
            "skuAttributes": [dict(attribute) for attribute in resolved_sku_attributes],
        })
        listing["images"].append({
            "imageKey": f"SKU-MAIN-{index:03d}",
            "productKey": product_key,
            "skuKey": sku_key,
            "role": "sku_main",
            "localPath": item["skuImage"],
            "order": 1,
            "mediaCategoryId": reference["media"]["defaultCategoryId"],
            "upload": {"status": "pending", "imgId": "", "imgUrl": ""},
        })
    for index, path in enumerate(source["mainImages"], start=1):
        listing["images"].append({
            "imageKey": f"SPU-MAIN-{index:02d}", "productKey": product_key, "skuKey": "",
            "role": "spu_main", "localPath": path, "order": index,
            "mediaCategoryId": reference["media"]["defaultCategoryId"],
            "upload": {"status": "pending", "imgId": "", "imgUrl": ""},
        })
    if str(listing["product"].get("styleImageMode") or "shared") == "series":
        scenes_by_style: dict[str, dict[str, str]] = {}
        for item in prepared:
            if item["style"] not in scenes_by_style:
                scenes_by_style[item["style"]] = (
                    source.get("styleSceneImagesBySku", {}).get(item["sourceSku"]) or {}
                )
        for style_order, style in enumerate(style_names, start=1):
            for position, path in sorted(
                ((int(pos), path) for pos, path in scenes_by_style.get(style, {}).items()),
                key=lambda pair: pair[0],
            ):
                listing["images"].append({
                    "imageKey": f"STYLE-{style_order:02d}-POS{position:02d}",
                    "productKey": product_key, "skuKey": "", "role": "style_scene",
                    "localPath": path, "order": position,
                    "mediaCategoryId": reference["media"]["defaultCategoryId"],
                    "upload": {"status": "pending", "imgId": "", "imgUrl": ""},
                })
    for index, path in enumerate(source["detailImages"], start=1):
        listing["images"].append({
            "imageKey": f"DETAIL-{index:03d}", "productKey": product_key, "skuKey": "",
            "role": "detail", "localPath": path, "order": index,
            "mediaCategoryId": reference["media"]["defaultCategoryId"],
            "upload": {"status": "pending", "imgId": "", "imgUrl": ""},
        })
    return listing


def build_audit_workbook(
    source: dict[str, Any], listing: dict[str, Any], destination: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU明细"
    sheet.append([f"商品标题：{listing['product']['title']}"])
    sheet.append([
        "序号", "SKU图", "款式名称", "规格尺寸", "商家SKU", "同行售价",
        "低于同行", "国补优惠", "活动优惠", "优惠券", "京东价", "预计到手价",
        "库存覆盖", "装裱", "题材", "是否启用", "源SKU", "校验结果",
    ])
    source_by_sku = {item["sourceSku"]: item for item in source["variants"]}
    embedded_images: set[str] = set()
    image_buffers: list[io.BytesIO] = []
    last_row = len(listing["skus"]) + 2
    for row_index, sku in enumerate(listing["skus"], start=3):
        source_row = source_by_sku[sku["sourceSku"]]
        sheet.append([
            "=ROW()-2", "", sku["styleAlias"], sku["sizeAlias"], sku.get("vendorSku") or "", source_row["sourcePrice"], 0.10,
            0.15, 0.15, 0,
            f'=IF(F{row_index}="","",ROUND((F{row_index}-G{row_index}+J{row_index})/((1-H{row_index})*(1-I{row_index})),2))',
            f'=IF(K{row_index}="","",ROUND(K{row_index}*(1-H{row_index})*(1-I{row_index})-J{row_index},2))',
            sku["stock"], sku["mountingLabel"], sku["subjectLabel"], "是", sku["sourceSku"],
            f'=IF(OR(C{row_index}="",D{row_index}="",E{row_index}="",F{row_index}="",G{row_index}="",H{row_index}="",I{row_index}="",J{row_index}="",K{row_index}="",L{row_index}="",M{row_index}="",N{row_index}="",O{row_index}="",P{row_index}=""),"缺少必填",IF(COUNTIF($E$3:$E${last_row},E{row_index})>1,"SKU编码重复","通过"))',
        ])
        sheet.cell(row_index, 5).number_format = "@"
        sheet.cell(row_index, 17).number_format = "@"
        image_path = source_row["skuImage"]
        if image_path not in embedded_images:
            image_buffer = io.BytesIO((destination.parent / image_path).read_bytes())
            image_buffers.append(image_buffer)
            image = ExcelImage(image_buffer)
            image.width = 72
            image.height = 72
            sheet.add_image(image, f"B{row_index}")
            embedded_images.add(image_path)
        sheet.row_dimensions[row_index].height = 58
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:R{last_row}"
    for column, width in {
        "A": 8, "B": 14, "C": 24, "D": 22, "E": 18, "F": 12, "K": 12,
        "L": 12, "M": 12, "N": 12, "O": 12, "P": 12, "Q": 18, "R": 14,
    }.items():
        sheet.column_dimensions[column].width = width

    summary = workbook.create_sheet("商品信息")
    summary.append(["字段", "值"])
    for row in [
        ("源链接", source["sourceUrl"]), ("源SKU", source["sourceSku"]),
        ("源SPU", source["sourceSpu"]), ("标题", listing["product"]["title"]),
        ("末级类目", f"{listing['category']['name']} ({listing['category']['id']})"),
        ("品牌", "无品牌"), ("工艺", "印刷"),
        ("运费模板", listing["product"]["freightTemplateLabel"]),
        ("主图模式", "多套主图"),
        ("7天无理由", listing["product"]["sevenDayReturnLabel"]),
        ("价格说明", "京东价按默认15%国补、15%直降和0元券反推；上架后由用户批量报名活动"),
        ("主图视频", (listing.get("video") or {}).get("localPath") or "无"),
        ("商家编码前缀", listing["product"]["vendorSkuPrefix"]),
        ("款式起始号", listing["product"]["vendorStyleStart"]),
        ("SKU数", len(listing["skus"])),
    ]:
        summary.append(row)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 90
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    for image_buffer in image_buffers:
        image_buffer.close()


def collect_to_folder(
    url: str, output_root: Path, rules_path: Path,
    reference_path: Path, style_start: int,
) -> dict[str, Any]:
    client = HttpClient()
    source = collect_source(url, client)
    folder = materialize(source, output_root, client)
    source = read_json(folder / "manifest.json")
    listing = build_listing(source, folder, read_json(rules_path), read_json(reference_path), style_start)
    write_json(folder / "listing.json", listing)
    workbook_path = folder / f"{folder.name}.xlsx"
    build_audit_workbook(source, listing, workbook_path)
    return {
        "ok": True,
        "folder": str(folder),
        "manifest": str(folder / "manifest.json"),
        "listing": str(folder / "listing.json"),
        "workbook": str(workbook_path),
        "sourceSku": source["sourceSku"],
        "sourceSpu": source["sourceSpu"],
        "skuCount": len(listing["skus"]),
        "mainImageCount": len(source["mainImages"]),
        "detailImageCount": len(source["detailImages"]),
        "vendorSkuPrefix": listing["product"]["vendorSkuPrefix"],
        "vendorStyleStart": style_start,
        "videoHandoff": video_handoff(folder, listing),
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    collect = commands.add_parser("collect")
    collect.add_argument("--url", required=True)
    collect.add_argument("--output", type=Path, default=ROOT / "outputs" / "migrations")
    collect.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    collect.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    collect.add_argument("--style-start", type=int, default=10000)
    args = parser.parse_args()
    try:
        result = collect_to_folder(args.url, args.output, args.rules, args.reference, args.style_start)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (MoverError, OSError, ValueError, requests.RequestException) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
